"""Рендеринг исходников корпуса в настоящие форматы: DOCX, PDF, XLSX, EML, TXT.

Корпус из .md-файлов сделал бы проект проще, чем он есть: разбор реальных форматов и восстановление
нумерации пунктов — половина работы приёма (docs/spec/06, §6). Поэтому исходник один, а на выходе —
тот формат, в котором документ жил бы на предприятии.

Поддерживаемое подмножество markdown: заголовки `##`…`#####` (это разделы и пункты), абзацы, списки `-`,
таблицы в виде `| … | … |`, цитаты `>`.
"""

from __future__ import annotations

import hashlib
import random
import re
from datetime import date, datetime
from email.message import EmailMessage
from email.utils import format_datetime
from functools import partial
from pathlib import Path
from typing import Any

import yaml

from rodos.corpus.source import SourceDoc, corpus_root

FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")
MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "eml": "message/rfc822",
    "txt": "text/plain",
    "xml": "application/xml",
}
TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
SEPARATOR_ROW = re.compile(r"^[\s|:-]+$")

# Перегенерация обязана давать побайтово те же файлы: иначе хеш документа меняется сам по себе, а
# сравнивать прогоны оценки между собой становится нельзя (docs/spec/06, §3). Поэтому всё, что
# библиотеки норовят взять из системных часов — метаданные документа и отметки времени внутри zip, —
# фиксируется на «сегодня» корпуса.
BUILD_TIME = datetime(2026, 9, 22, 12, 0, 0)
ZIP_TIME = (BUILD_TIME.year, BUILD_TIME.month, BUILD_TIME.day, BUILD_TIME.hour, BUILD_TIME.minute,
            BUILD_TIME.second)
AUTHOR = "ГК «Родос-Деталь»"


def _freeze_zip(path: Path) -> None:
    """Переписывает DOCX/XLSX: отметки времени записей архива и даты в свойствах документа.

    Свойства приходится править уже в архиве: openpyxl ставит `modified` в момент сохранения,
    перекрывая то, что задано в объекте книги.
    """
    import zipfile

    stamp = f"{BUILD_TIME:%Y-%m-%dT%H:%M:%S}Z".encode()
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, payload in entries:
            if info.filename == "docProps/core.xml":
                payload = re.sub(rb">[0-9T:\-]+Z<", b">" + stamp + b"<", payload)
            frozen = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            frozen.compress_type = info.compress_type
            frozen.external_attr = info.external_attr
            archive.writestr(frozen, payload)


def _blocks(body: str) -> list[tuple[str, Any]]:
    """Тело документа → последовательность блоков (вид, содержимое)."""
    blocks: list[tuple[str, Any]] = []
    lines = body.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append(("heading", (level, stripped[level:].strip())))
            index += 1
        elif TABLE_ROW.match(line):
            rows: list[list[str]] = []
            while index < len(lines) and TABLE_ROW.match(lines[index]):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not SEPARATOR_ROW.match(lines[index].strip().strip("|")):
                    rows.append(cells)
                index += 1
            blocks.append(("table", rows))
        elif stripped.startswith("- "):
            items: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("- "):
                items.append(lines[index].strip()[2:].strip())
                index += 1
            blocks.append(("bullets", items))
        elif stripped.startswith("> "):
            blocks.append(("note", stripped[2:].strip()))
            index += 1
        else:
            paragraph = [stripped]
            index += 1
            while index < len(lines) and lines[index].strip() and not re.match(
                    r"^\s*(#|-\s|>\s|\|)", lines[index]):
                paragraph.append(lines[index].strip())
                index += 1
            blocks.append(("paragraph", " ".join(paragraph)))
    return blocks


def _plain(text: str) -> str:
    """Убирает markdown-разметку выделения, оставляя текст."""
    return re.sub(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`", lambda m: m.group(1) or m.group(2) or m.group(3), text)


def _header_lines(doc: SourceDoc) -> list[str]:
    """Шапка документа: утверждающая подпись, реквизиты, статус."""
    card = doc.card
    lines = []
    if approver := card.get("approved_by"):
        lines.append(f"УТВЕРЖДАЮ: {approver}")
    if approved := card.get("approved_at"):
        lines.append(f"Дата утверждения: {approved:%d.%m.%Y}")
    if version := card.get("version"):
        lines.append(f"Редакция: {version}")
    if effective := card.get("effective_from"):
        lines.append(f"Вводится в действие с {effective:%d.%m.%Y}")
    if supersedes := card.get("supersedes"):
        lines.append(f"Взамен: {supersedes}")
    if card.get("confidentiality") == "dsp":
        lines.append("Для служебного пользования")
    if card.get("status") == "superseded":
        lines.append(f"Документ заменён: {card.get('superseded_by')}")
    if card.get("status") == "retracted":
        lines.append("Документ отменён")
    return lines


def render_docx(doc: SourceDoc, target: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    for line in _header_lines(doc):
        paragraph = document.add_paragraph(line)
        paragraph.runs[0].italic = True
    document.add_heading(str(doc.card["title"]), level=0)

    for kind, payload in _blocks(doc.body):
        if kind == "heading":
            level, text = payload
            document.add_heading(_plain(text), level=min(level, 4))
        elif kind == "paragraph":
            document.add_paragraph(_plain(payload))
        elif kind == "bullets":
            for item in payload:
                document.add_paragraph(_plain(item), style="List Bullet")
        elif kind == "note":
            paragraph = document.add_paragraph(_plain(payload))
            paragraph.runs[0].italic = True
        elif kind == "table":
            rows: list[list[str]] = payload
            table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for cell_index, value in enumerate(row):
                    table.cell(row_index, cell_index).text = _plain(value)
    core = document.core_properties
    core.author = core.last_modified_by = AUTHOR
    core.created = core.modified = BUILD_TIME
    core.revision = 1
    document.save(str(target))
    _freeze_zip(target)


def render_pdf(doc: SourceDoc, target: Path) -> None:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF(format="A4")
    pdf.set_creation_date(BUILD_TIME)
    pdf.set_author(AUTHOR)
    # fpdf после multi_cell оставляет курсор справа; всё, что пишем, начинается от левого поля.
    cell = partial(pdf.multi_cell, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.add_font("DejaVu", "", str(FONT_DIR / "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", str(FONT_DIR / "DejaVuSans-Bold.ttf"))
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("DejaVu", "", 9)
    for line in _header_lines(doc):
        cell(0, 5, line)
    pdf.ln(2)
    pdf.set_font("DejaVu", "B", 13)
    cell(0, 7, str(doc.card["title"]))
    pdf.ln(2)

    for kind, payload in _blocks(doc.body):
        if kind == "heading":
            level, text = payload
            pdf.set_font("DejaVu", "B", max(12 - level, 9))
            cell(0, 6, _plain(text))
        elif kind == "paragraph":
            pdf.set_font("DejaVu", "", 10)
            cell(0, 5, _plain(payload))
        elif kind == "bullets":
            pdf.set_font("DejaVu", "", 10)
            for item in payload:
                cell(0, 5, f"• {_plain(item)}")
        elif kind == "note":
            pdf.set_font("DejaVu", "", 9)
            cell(0, 5, _plain(payload))
        elif kind == "table":
            rows: list[list[str]] = payload
            pdf.set_font("DejaVu", "", 8)
            columns = max(len(row) for row in rows)
            available = pdf.w - 2 * pdf.l_margin
            # Ширина колонки пропорциональна длине её самого длинного значения, но не уже 18 мм.
            longest = [max((len(_plain(row[index])) for row in rows if index < len(row)), default=1)
                       for index in range(columns)]
            widths = [max(18.0, available * value / sum(longest)) for value in longest]
            scale = available / sum(widths)
            widths = [width * scale for width in widths]
            with pdf.table(col_widths=widths, line_height=5, first_row_as_headings=len(rows) > 1,
                           text_align="LEFT") as table:
                for row in rows:
                    table_row = table.row()
                    for index in range(columns):
                        table_row.cell(_plain(row[index]) if index < len(row) else "")
        pdf.ln(1)
    pdf.output(str(target))


def render_xlsx(doc: SourceDoc, target: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    assert doc.table is not None, f"{doc.path}: нет описания таблицы"
    workbook = Workbook()
    sheets = doc.table if isinstance(doc.table, list) else [doc.table]
    default = workbook.active
    assert default is not None
    workbook.remove(default)
    for sheet_spec in sheets:
        sheet = workbook.create_sheet(str(sheet_spec.get("name", "Лист"))[:31])
        row_index = 1
        if title := sheet_spec.get("title"):
            sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=12)
            row_index = 3
        columns: list[str] = sheet_spec["columns"]
        for column_index, name in enumerate(columns, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=name)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for record in sheet_spec["rows"]:
            row_index += 1
            values = record if isinstance(record, list) else [record.get(name) for name in columns]
            for column_index, value in enumerate(values, start=1):
                sheet.cell(row=row_index, column=column_index, value=value)
        for column_index, name in enumerate(columns, start=1):
            longest = max([len(str(name))] + [
                len(str((record if isinstance(record, list) else
                         [record.get(item) for item in columns])[column_index - 1] or ""))
                for record in sheet_spec["rows"]])
            sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].width = min(
                max(12, longest + 2), 60)
    workbook.properties.creator = workbook.properties.lastModifiedBy = AUTHOR
    workbook.properties.created = workbook.properties.modified = BUILD_TIME
    workbook.save(str(target))
    _freeze_zip(target)


def render_eml(doc: SourceDoc, target: Path) -> None:
    card = doc.card
    message = EmailMessage()
    message["From"] = card["mail_from"]
    message["To"] = card["mail_to"]
    message["Subject"] = card.get("mail_subject", card["title"])
    sent: date = card.get("mail_date") or card.get("approved_at") or date(2026, 9, 1)
    message["Date"] = format_datetime(datetime(sent.year, sent.month, sent.day, 10, 30))
    message["Message-ID"] = f"<{doc.doc_id}@rodos-detal.example.com>"
    if reply_to := card.get("mail_in_reply_to"):
        message["In-Reply-To"] = f"<{reply_to}@rodos-detal.example.com>"
    message.set_content("\n".join(_plain(line) for line in doc.body.splitlines()))
    target.write_bytes(message.as_bytes())


def render_txt(doc: SourceDoc, target: Path) -> None:
    lines = [*_header_lines(doc), "", str(doc.card["title"]), ""]
    lines += [_plain(line) for line in doc.body.splitlines()]
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")


def render_pdf_scan(doc: SourceDoc, target: Path) -> None:
    """PDF без текстового слоя: страница — одна растровая картинка, как со сканера.

    Нужен, чтобы доказать честный путь «карточка есть, содержимого нет»: приём обязан завести документ,
    не выдумывая текст, и не пускать его в ответы (docs/spec/03, §3).
    """
    from fpdf import FPDF
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1240, 1754  # A4 при 150 dpi
    page = Image.new("L", (width, height), 244)
    draw = ImageDraw.Draw(page)
    title_font = ImageFont.truetype(str(FONT_DIR / "DejaVuSans-Bold.ttf"), 30)
    body_font = ImageFont.truetype(str(FONT_DIR / "DejaVuSans.ttf"), 23)

    y = 150
    draw.text((110, y), str(doc.card["title"]), font=title_font, fill=40)
    y += 70
    for kind, payload in _blocks(doc.body):
        text = _plain(payload if isinstance(payload, str) else str(payload[1]))
        font = title_font if kind == "heading" else body_font
        for line in _wrap(text, 64 if kind == "heading" else 78):
            draw.text((110, y), line, font=font, fill=55)
            y += 38 if kind == "heading" else 32
        y += 12
        if y > height - 150:
            break

    # Признаки скана: лёгкий перекос и шум — иначе документ слишком чистый, чтобы быть похожим.
    # Шум берётся из генератора, засеянного идентификатором документа: `Image.effect_noise` засевается
    # системной энтропией и ломает побайтовую воспроизводимость сборки (docs/spec/06, §3).
    page = page.rotate(-0.45, resample=Image.Resampling.BILINEAR, fillcolor=244)
    grain = random.Random(doc.doc_id).randbytes(width * height)
    noise = Image.frombytes("L", (width, height), grain).point(lambda value: value // 8 + 240)
    page = Image.blend(page.convert("L"), noise.convert("L"), 0.10)

    raster = target.with_suffix(".scan.png")
    page.save(raster, format="PNG")
    pdf = FPDF(format="A4")
    pdf.set_creation_date(BUILD_TIME)
    pdf.set_author(AUTHOR)
    pdf.set_auto_page_break(auto=False)
    pdf.add_page()
    pdf.image(str(raster), x=0, y=0, w=210, h=297)
    pdf.output(str(target))
    raster.unlink()


def _wrap(text: str, width: int) -> list[str]:
    words, lines, current = text.split(), [], ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) > width and current:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines or [""]


def render_xml(doc: SourceDoc, target: Path) -> None:
    """Первичка ЭДО: тело исходника уже XML, пишем как есть.

    Извлечение из XML — отдельная ветка приёма: карточка почти целиком структурна, и проверять её
    нужно на настоящем дереве, а не на пересказе (docs/spec/06, §6).
    """
    from xml.etree import ElementTree

    payload = doc.body.strip()
    ElementTree.fromstring(payload)  # исходник обязан быть разбираемым XML
    target.write_text(payload + "\n", encoding="utf-8")


RENDERERS = {"docx": render_docx, "pdf": render_pdf, "xlsx": render_xlsx, "eml": render_eml,
             "txt": render_txt, "xml": render_xml, "pdf_scan": render_pdf_scan}
# Скан — это тоже .pdf: формат исходника говорит, как рендерить, а не как называется файл.
EXTENSIONS = {"pdf_scan": "pdf"}


def _repo_path(path: Path, base: Path) -> str:
    """Путь относительно базы, если он внутри неё, иначе абсолютный.

    Карточка должна читаться и тогда, когда рендер идёт мимо репозитория, а не падать на вычислении
    относительного пути.
    """
    try:
        return str(path.relative_to(base))
    except ValueError:
        return str(path)


def render(doc: SourceDoc, root: Path | None = None) -> tuple[Path, dict[str, Any]]:
    """Рендерит документ и пишет эталонную карточку; возвращает путь и карточку."""
    base = root or corpus_root()
    target_dir = base / "rendered" / doc.space
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{doc.doc_id}.{EXTENSIONS.get(doc.fmt, doc.fmt)}"
    if doc.fmt not in RENDERERS:
        raise ValueError(f"{doc.path}: формат {doc.fmt} не поддерживается")
    RENDERERS[doc.fmt](doc, target)

    card = dict(doc.card)
    payload = target.read_bytes()
    card.update({
        # Исходник — это происхождение документа: он лежит в репозитории независимо от того, куда
        # сейчас рендерим. Сборка во временный каталог (тесты) не должна менять эту запись.
        "source": _repo_path(doc.path, corpus_root().parent),
        "rendered": _repo_path(target, base.parent),
        "mime": MIME[EXTENSIONS.get(doc.fmt, doc.fmt)],
        "bytes": len(payload),
        "sha256": hashlib.sha256(payload).hexdigest(),
    })
    cards_dir = base / "cards"
    cards_dir.mkdir(parents=True, exist_ok=True)
    (cards_dir / f"{doc.doc_id}.yaml").write_text(
        yaml.dump(card, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return target, card
